"""
Delvix AI — Melbourne SMB Lead Scraper
Run: python main.py
"""

import sys
from dotenv import load_dotenv

load_dotenv()

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from config import OUTPUT_FILE
from scraper import run_scrape
from enricher import enrich_leads


COLUMNS = [
    ("Business Name",     "business_name",    35),
    ("Category",          "category",         28),
    ("Suburb",            "suburb",           18),
    ("Address",           "address",          40),
    ("Phone",             "phone",            18),
    ("Website",           "website",          35),
    ("Email (Maps)",      "email_maps",       30),
    ("Emails (Website)",  "emails_website",   40),
    ("Rating",            "rating",           10),
    ("Reviews",           "reviews",          10),
    ("Google Maps URL",   "google_maps_url",  50),
]

HEADER_FILL   = PatternFill("solid", fgColor="1B2A4A")
HEADER_FONT   = Font(bold=True, color="FFFFFF", name="Calibri", size=11)
ALT_ROW_FILL  = PatternFill("solid", fgColor="EEF2F7")
LINK_FONT     = Font(color="1155CC", underline="single", name="Calibri", size=10)
BODY_FONT     = Font(name="Calibri", size=10)
THIN_BORDER   = Border(
    bottom=Side(style="thin", color="D0D7E3"),
)


def _write_excel(leads: list[dict], path: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Melbourne Leads"
    ws.freeze_panes = "A2"

    # Header row
    for col_idx, (header, _, width) in enumerate(COLUMNS, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.row_dimensions[1].height = 22

    # Data rows
    for row_idx, lead in enumerate(leads, 2):
        fill = ALT_ROW_FILL if row_idx % 2 == 0 else None
        for col_idx, (_, key, _) in enumerate(COLUMNS, 1):
            value = lead.get(key, "")
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.border = THIN_BORDER
            cell.alignment = Alignment(vertical="center", wrap_text=False)

            # Hyperlink columns
            if key == "website" and value and value.startswith("http"):
                cell.hyperlink = value
                cell.font = LINK_FONT
            elif key == "google_maps_url" and value and value.startswith("http"):
                cell.hyperlink = value
                cell.value = "Open Map"
                cell.font = LINK_FONT
            else:
                cell.font = BODY_FONT

            if fill:
                cell.fill = fill

    # Auto-filter
    ws.auto_filter.ref = ws.dimensions

    # Summary sheet
    ws2 = wb.create_sheet("Summary")
    ws2["A1"] = "Category"
    ws2["B1"] = "Count"
    ws2["A1"].font = HEADER_FONT
    ws2["A1"].fill = HEADER_FILL
    ws2["B1"].font = HEADER_FONT
    ws2["B1"].fill = HEADER_FILL
    ws2.column_dimensions["A"].width = 35
    ws2.column_dimensions["B"].width = 12

    counts: dict[str, int] = {}
    for lead in leads:
        counts[lead["category"]] = counts.get(lead["category"], 0) + 1

    for r, (cat, count) in enumerate(sorted(counts.items()), 2):
        ws2.cell(row=r, column=1, value=cat).font = BODY_FONT
        ws2.cell(row=r, column=2, value=count).font = BODY_FONT

    total_row = len(counts) + 2
    ws2.cell(row=total_row, column=1, value="TOTAL").font = Font(bold=True, name="Calibri", size=10)
    ws2.cell(row=total_row, column=2, value=len(leads)).font = Font(bold=True, name="Calibri", size=10)

    wb.save(path)


def main():
    print("=" * 60)
    print("  Delvix AI — Melbourne SMB Lead Scraper")
    print("=" * 60)

    print("\n[1/3] Scraping Google Maps via Outscraper...")
    leads = run_scrape()
    print(f"\n  -> {len(leads)} unique businesses found.")

    print("\n[2/3] Enriching with website emails...")
    enrich_leads(leads)

    print(f"\n[3/3] Writing spreadsheet: {OUTPUT_FILE}")
    _write_excel(leads, OUTPUT_FILE)

    with_email = sum(
        1 for l in leads
        if l.get("email_maps") or l.get("emails_website")
    )
    print(f"\nDone. {len(leads)} leads total, {with_email} with at least one email.")
    print(f"Output: {OUTPUT_FILE}")


if __name__ == "__main__":
    main()
